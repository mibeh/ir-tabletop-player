from __future__ import annotations

import sys
from pathlib import Path

import pytest


def _minimal_findings() -> dict:
    return {
        "exercise_overview": {
            "name": "Operation Paper Trail",
            "dates": "2026-05-29",
            "scope": "A TTX of ~2.5 hours for an internal organization.",
            "mission_areas": ["Protection", "Response"],
            "core_capabilities": ["Cybersecurity", "Operational Coordination"],
            "objectives": [
                "Exercise the BEC playbook end-to-end.",
                "Practice cross-team escalation.",
            ],
            "threat_or_hazard": "Cyber Incident — Business Email Compromise",
            "scenario_description": "An AiTM phishing attack against Finance.",
            "modules": [
                {"number": 1, "title": "Detection"},
                {"number": 2, "title": "Containment"},
            ],
            "sponsor": "IR Team",
            "participating_organizations": "SOC, IR, Finance, Legal.",
            "points_of_contact": "Facilitator: Jane Doe, jane@example.org",
            "caveat": "Fictional Training Exercise — Not a Real Incident",
        },
        "executive_summary": (
            "The exercise tested the organization's response to a finance-targeted "
            "business email compromise across detection, containment, and "
            "recovery.\n\n"
            "Participants demonstrated strong within-SLA triage and early legal "
            "engagement. The most pressing area for improvement is the absence "
            "of a session-revocation step in the BEC playbook."
        ),
        "objective_capability_alignment": [
            {"objective": "Exercise the BEC playbook end-to-end.",
             "core_capabilities": ["Cybersecurity"]},
            {"objective": "Practice cross-team escalation.",
             "core_capabilities": ["Operational Coordination"]},
        ],
        "objective_analyses": [
            {
                "number": 1,
                "objective": "Exercise the BEC playbook end-to-end.",
                "core_capabilities": ["Cybersecurity"],
                "strengths": ["The SOC triaged the Scanner alert within SLA."],
                "areas_for_improvement": [
                    {
                        "observation": "Session revocation was not the first instinct.",
                        "priority": "High",
                        "references": ["BEC Playbook v3"],
                        "analysis": "The team reset the password before revoking sessions.",
                        "recommendation": "Add an explicit session-revocation step.",
                    }
                ],
            },
            {
                "number": 2,
                "objective": "Practice cross-team escalation.",
                "core_capabilities": ["Operational Coordination"],
                "strengths": ["Legal was engaged early."],
                "areas_for_improvement": [],
            },
        ],
        "improvement_plan": [
            {
                "exercise_objective": "Exercise the BEC playbook end-to-end.",
                "core_capability": "Cybersecurity",
                "issue": "No session-revocation step in the playbook.",
                "priority": "High",
                "corrective_actions": ["Update the BEC playbook."],
                "capability_element": "Planning",
                "responsible_organization": "IR",
                "poc": "Jane Doe",
                "start_date": "2026-06-15",
                "completion_date": "2026-07-15",
            }
        ],
        "participants": [
            {
                "name": "Alex Chen",
                "position": "SOC Analyst",
                "exercise_role": "Participant",
                "department": "SOC",
            },
            {
                "name": "Jordan Lee",
                "position": "IR Lead",
                "exercise_role": "Facilitator",
                "department": "IR",
            },
        ],
    }


@pytest.fixture
def findings() -> dict:
    return _minimal_findings()


def _all_text(doc) -> str:
    chunks = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            chunks.extend(cell.text for cell in row.cells)
    return "\n".join(chunks)


def test_slugify(render_aar):
    assert render_aar._slugify("Operation Paper Trail!") == "operation-paper-trail"
    assert render_aar._slugify("   ") == ""
    assert len(render_aar._slugify("a" * 70)) == 60


def test_renders_docx_file(render_aar, findings, tmp_path):
    out = render_aar.render_aar(findings, tmp_path)
    assert out.exists()
    assert out.suffix == ".docx"
    assert out.parent == tmp_path


def test_title_page_has_exercise_name_and_title(render_aar, findings, tmp_path):
    from docx import Document

    out = render_aar.render_aar(findings, tmp_path)
    paragraphs = [p.text for p in Document(str(out)).paragraphs]
    title_page = paragraphs[:paragraphs.index("EXERCISE OVERVIEW")]
    assert "Operation Paper Trail" in title_page
    assert any("After-Action Report / Improvement Plan" in p for p in title_page)


def test_title_page_has_no_organization_logo(render_aar, findings, tmp_path):
    from docx import Document

    out = render_aar.render_aar(findings, tmp_path)
    doc = Document(str(out))
    assert not doc.inline_shapes


def test_overview_table_has_expected_labels(render_aar, findings, tmp_path):
    from docx import Document

    out = render_aar.render_aar(findings, tmp_path)
    text = _all_text(Document(str(out)))
    for label in [
        "Exercise Name",
        "Exercise Dates",
        "Scope",
        "Mission Area(s)",
        "Core Capabilities",
        "Objectives",
        "Threat or Hazard",
        "Scenario",
        "Sponsor",
        "Participating Organizations",
        "Points of Contact",
    ]:
        assert label in text, f"missing overview label: {label}"
    assert "EXERCISE OVERVIEW" in text
    assert "Cybersecurity" in text
    assert "Detection" in text  # module title appears in Scenario row


def test_objective_analysis_sections(render_aar, findings, tmp_path):
    from docx import Document

    out = render_aar.render_aar(findings, tmp_path)
    text = _all_text(Document(str(out)))
    assert "ANALYSIS OF EXERCISE OBJECTIVES" in text
    assert "Objective 1: Exercise the BEC playbook end-to-end." in text
    assert "Objective 2: Practice cross-team escalation." in text
    assert "Strengths" in text
    assert "Areas for Improvement" in text
    assert "Session revocation was not the first instinct." in text
    assert "BEC Playbook v3" in text  # reference rendered
    assert "Add an explicit session-revocation step." in text  # recommendation
    assert "The SOC triaged the Scanner alert within SLA." in text  # strength
    assert "No areas for improvement recorded for this objective." in text


def test_area_for_improvement_shows_priority(render_aar, findings, tmp_path):
    from docx import Document

    out = render_aar.render_aar(findings, tmp_path)
    paragraphs = [p.text for p in Document(str(out)).paragraphs]
    start = paragraphs.index("ANALYSIS OF EXERCISE OBJECTIVES")
    end = paragraphs.index("APPENDIX A: IMPROVEMENT PLAN")
    assert "Priority: High" in paragraphs[start:end]


def test_executive_summary_section(render_aar, findings, tmp_path):
    from docx import Document

    out = render_aar.render_aar(findings, tmp_path)
    doc = Document(str(out))
    text = _all_text(doc)
    assert "EXECUTIVE SUMMARY" in text
    assert "finance-targeted" in text
    assert "session-revocation step in the BEC playbook" in text
    # The summary must appear after the overview and before the objective analysis.
    headings = [p.text for p in doc.paragraphs]
    assert headings.index("EXECUTIVE SUMMARY") > headings.index("EXERCISE OVERVIEW")
    assert headings.index("EXECUTIVE SUMMARY") < headings.index(
        "ANALYSIS OF EXERCISE OBJECTIVES"
    )


def test_executive_summary_defaults_when_absent(render_aar, findings, tmp_path):
    from docx import Document

    del findings["executive_summary"]
    out = render_aar.render_aar(findings, tmp_path)
    text = _all_text(Document(str(out)))
    assert "EXECUTIVE SUMMARY" in text
    assert "No executive summary was provided" in text


def test_all_tables_autofit_contents(render_aar, findings, tmp_path):
    from docx import Document
    from docx.oxml.ns import qn

    out = render_aar.render_aar(findings, tmp_path)
    doc = Document(str(out))
    assert doc.tables, "expected at least one table"
    for table in doc.tables:
        tbl_pr = table._tbl.tblPr
        layout = tbl_pr.find(qn("w:tblLayout"))
        assert layout is not None and layout.get(qn("w:type")) == "autofit"
        tbl_w = tbl_pr.find(qn("w:tblW"))
        assert tbl_w is not None and tbl_w.get(qn("w:type")) == "auto"


def test_improvement_plan_matrix(render_aar, findings, tmp_path):
    from docx import Document

    out = render_aar.render_aar(findings, tmp_path)
    doc = Document(str(out))
    text = _all_text(doc)
    assert "APPENDIX A: IMPROVEMENT PLAN" in text
    for header in [
        "Exercise Objective",
        "Core Capability",
        "Issue/Area for Improvement",
        "Priority",
        "Corrective Action",
        "Capability Element",
        "Primary Responsible Organization",
        "Organization POC",
        "Start Date",
        "Completion Date",
    ]:
        assert header in text, f"missing IP column: {header}"
    assert "Update the BEC playbook." in text
    assert "2026-07-15" in text
    assert "2026-06-15" in text  # start_date rendered
    ip_table = next(
        t for t in doc.tables if t.rows[0].cells[0].text == "Exercise Objective"
    )
    assert len(ip_table.columns) == 10
    priority_col = [c.text for c in ip_table.columns[3].cells]
    assert priority_col[0] == "Priority"
    assert "High" in priority_col[1]


def test_participants_appendix(render_aar, findings, tmp_path):
    from docx import Document

    out = render_aar.render_aar(findings, tmp_path)
    text = _all_text(Document(str(out)))
    assert "APPENDIX B: EXERCISE PARTICIPANTS" in text
    for header in ["Name", "Position", "Exercise Role", "Department"]:
        assert header in text, f"missing participant column: {header}"
    assert "Alex Chen" in text
    assert "SOC Analyst" in text  # position
    assert "Participant" in text  # exercise_role
    assert "Jordan Lee" in text
    assert "Facilitator" in text  # exercise_role


def test_header_and_footer_carry_classification(render_aar, findings, tmp_path):
    from docx import Document

    out = render_aar.render_aar(findings, tmp_path)
    doc = Document(str(out))
    header_text = "\n".join(p.text for p in doc.sections[0].header.paragraphs)
    footer_text = "\n".join(p.text for p in doc.sections[0].footer.paragraphs)
    assert "TLP:AMBER+STRICT - CONFIDENTIAL" in header_text
    assert "TLP:AMBER+STRICT - CONFIDENTIAL" in footer_text


def test_cli_main_writes_docx(render_aar, tmp_path, monkeypatch, capsys):
    findings_path = Path(__file__).resolve().parent / "fixtures" / "sample_findings.json"
    monkeypatch.setattr(
        sys, "argv",
        ["render_aar.py", str(findings_path), "--output-dir", str(tmp_path)],
    )
    render_aar.main()
    captured = capsys.readouterr().out
    assert "Wrote" in captured
    written = list(tmp_path.glob("aar-*.docx"))
    assert len(written) == 1

    from docx import Document

    Document(str(written[0]))  # opens without error -> well-formed .docx


def test_sponsor_defaults_in_improvement_plan_narrative(render_aar, findings, tmp_path):
    from docx import Document

    ov = findings["exercise_overview"]
    del ov["sponsor"]
    out = render_aar.render_aar(findings, tmp_path)
    text = _all_text(Document(str(out)))
    # With no sponsor in the overview, the Improvement Plan narrative falls back
    # to the default sponsor.
    assert "This Improvement Plan was developed for IR Team" in text
